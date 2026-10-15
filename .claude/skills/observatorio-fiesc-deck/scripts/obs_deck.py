"""
obs_deck — decks no padrão Observatório FIESC (modelo "Panorama Setor da Madeira", out/2026).

Clona os slides-base de assets/template.pptx (capa, seção, conteúdo, síntese, encerramento)
e preenche texto + gráficos NATIVOS do PowerPoint (dados em Excel embutido).

    from obs_deck import ObsDeck
    deck = ObsDeck()
    deck.capa("Setor da Madeira", "Panorama setorial", "Outubro de 2026")
    deck.secao("01", "Comércio", "exterior")
    deck.grafico(titulo=("Exportações do complexo florestal ", "recuam"), ...)
    deck.sintese(("Síntese ", "e leitura do período"), esquerda=[...], direita=[...])
    deck.encerramento()
    deck.salvar("reports/presentations/[OBS FIESC] Meu estudo.pptx")

Texto com destaque: use **trecho** → Montserrat SemiBold azul (#0240CA).
"""
import copy
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

TEMPLATE = Path(__file__).resolve().parents[1] / "assets" / "template.pptx"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"

# ── Paleta do deck final ──────────────────────────────────────────────────────
INK = "1E1E29"        # texto
BLUE = "0240CA"       # destaque / ano corrente
BLUE_L = "7F9BE6"     # ano anterior
GRAY = "B9BEC8"       # dois anos antes
GRID = "E4E7EC"       # eixos e linhas de tabela
MUTED = "6B7280"      # fonte
BG = "F6F6F0"         # fundo dos slides de conteúdo
SERIES = {1: [BLUE], 2: [BLUE_L, BLUE], 3: [GRAY, BLUE_L, BLUE]}

BASE = {"capa": 0, "secao": 1, "conteudo": 2, "sintese": 3, "encerramento": 4}


# ── Números pt-BR ─────────────────────────────────────────────────────────────
def num_br(x, nd=0):
    return f"{x:,.{nd}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def var_br(x, nd=1, suf=""):
    return ("+" if x > 0 else "") + num_br(x, nd) + suf


# ── Texto ─────────────────────────────────────────────────────────────────────
def _rpr(font, size, color, lang="pt-BR"):
    r = etree.SubElement(etree.Element("dummy"), qn("a:rPr"), lang=lang, sz=str(int(size * 100)), dirty="0")
    sf = etree.SubElement(r, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=color)
    etree.SubElement(r, qn("a:latin"), typeface=font)
    return r


def _partes(texto):
    """'a **b** c' -> [('a ', False), ('b', True), (' c', False)]"""
    out = []
    for i, p in enumerate(re.split(r"\*\*(.+?)\*\*", texto)):
        if p:
            out.append((p, i % 2 == 1))
    return out


def preencher(shape, paragrafos, size, normal=("Montserrat", INK), destaque=("Montserrat SemiBold", BLUE)):
    """Reescreve o text frame mantendo o pPr (alinhamento, entrelinha, espaçamento) do 1º parágrafo."""
    tf = shape.text_frame
    ps = tf._txBody.findall(qn("a:p"))
    ppr = ps[0].find(qn("a:pPr"))
    ppr = copy.deepcopy(ppr) if ppr is not None else None
    for p in ps:
        tf._txBody.remove(p)
    if isinstance(paragrafos, str):
        paragrafos = [paragrafos]
    for txt in paragrafos:
        p = etree.SubElement(tf._txBody, qn("a:p"))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        for parte, dest in _partes(txt):
            r = etree.SubElement(p, qn("a:r"))
            font, color = destaque if dest else normal
            r.append(_rpr(font, size, color))
            etree.SubElement(r, qn("a:t")).text = parte
    return shape


# ── Deck ──────────────────────────────────────────────────────────────────────
class ObsDeck:
    def __init__(self, template=TEMPLATE):
        self.prs = Presentation(str(template))
        self._base = list(self.prs.slides)

    # clona um slide-base (shapes, fundo e relações de imagem)
    def _clone(self, tipo):
        src = self._base[BASE[tipo]]
        dst = self.prs.slides.add_slide(src.slide_layout)
        for ph in list(dst.placeholders):
            ph._element.getparent().remove(ph._element)
        def _copia(el):
            novo = copy.deepcopy(el)
            for node in novo.iter():
                for att in (f"{{{R_NS}}}embed", f"{{{R_NS}}}link", f"{{{R_NS}}}id"):
                    rid = node.get(att)
                    if rid and rid in src.part.rels:
                        rel = src.part.rels[rid]
                        node.set(att, dst.part.relate_to(rel._target, rel.reltype))
            return novo

        if src._element.cSld.bg is not None:          # fundo pode ser imagem (capa/encerramento)
            dst._element.cSld.insert(0, _copia(src._element.cSld.bg))
        tree = dst.shapes._spTree
        for el in src.shapes._spTree.iterchildren():
            if el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
                tree.append(_copia(el))
        return dst

    @staticmethod
    def _shape(slide, nome):
        return next(s for s in slide.shapes if s.name == nome)

    def _titulo(self, slide, titulo, size=24):
        """titulo: str (todo escuro) ou (parte_escura, parte_azul)."""
        if isinstance(titulo, str):
            titulo = (titulo, "")
        escuro, azul = titulo
        txt = escuro + (f"**{azul}**" if azul else "")
        preencher(self._shape(slide, "titulo"), [txt], size,
                  normal=("Montserrat Medium", INK), destaque=("Montserrat Medium", BLUE))

    # ── slides ────────────────────────────────────────────────────────────────
    def capa(self, titulo, subtitulo, data):
        s = self._clone("capa")
        for nome, txt in (("capa_titulo", titulo), ("capa_subtitulo", subtitulo), ("capa_data", data)):
            sh = self._shape(s, nome)
            r = sh.text_frame.paragraphs[0].runs[0]
            r.text = txt
        return s

    def secao(self, numero, linha_fina, linha_forte):
        s = self._clone("secao")
        tf = self._shape(s, "secao_titulo").text_frame
        tf.paragraphs[0].runs[0].text = linha_fina
        tf.paragraphs[1].runs[0].text = linha_forte
        self._shape(s, "secao_numero").text_frame.paragraphs[0].runs[0].text = numero
        return s

    def grafico(self, titulo, grafico_titulo, unidade, texto, grafico, fonte,
                tabela=None, nota=None, tamanho_titulo=24, tamanho_texto=10.5):
        """Slide de conteúdo: gráfico à esquerda, análise curta à direita.

        grafico: dict(tipo='colunas'|'barras'|'linhas', categorias=[...],
                      series={'2024': [...], '2025': [...], '2026': [...]},  # última = ano corrente
                      formato='#,##0.00' (comex) | '#,##0' (empregos),
                      rotular='todas'|'ultima' (colunas/barras) | 'fim'|'ultima' (linhas),
                      eixo_min=None, eixo_max=None,   # só linhas
                      plot_x=0.33)                    # barras+tabela: fração reservada aos rótulos
        tabela:  dict(cabecalho=['Var. %', 'Var. US$ mi'], linhas=[['+7,1%', '+19,02'], ...])
                 — só com tipo='barras'; fica alinhada às barras, entre gráfico e texto.
        texto:   lista de parágrafos; **trecho** vira destaque azul.
        """
        s = self._clone("conteudo")
        self._titulo(s, titulo, tamanho_titulo)
        preencher(self._shape(s, "grafico_titulo"), [grafico_titulo], 11.5,
                  normal=("Montserrat SemiBold", BLUE))
        preencher(self._shape(s, "grafico_unidade"), [unidade], 9, normal=("Montserrat", BLUE))
        preencher(self._shape(s, "fonte"), [fonte], 8, normal=("Montserrat", MUTED))
        tx = self._shape(s, "texto")
        preencher(tx, texto, tamanho_texto)

        tipo = grafico["tipo"]
        com_tabela = tabela is not None and tipo == "barras"
        x, y, w, h = 0.35, 1.85, (4.30 if com_tabela else 5.95), 2.88
        chart = self._add_chart(s, grafico, x, y, w, h)
        if com_tabela:
            px = grafico.get("plot_x", 0.33)
            _manual_layout(chart, px, PLOT_Y, 0.97 - px - 0.08, PLOT_H)
            _tabela_alinhada(s, tabela, x + w + 0.05, y, h, len(grafico["categorias"]))
        tx.left, tx.top, tx.width = Inches(6.60), Inches(1.75), Inches(2.95)
        if nota:
            nb = s.shapes.add_textbox(Inches(6.20), Inches(4.45), Inches(3.40), Inches(0.9))
            nb.name = "nota"
            nb.text_frame.word_wrap = True
            preencher(nb, nota if isinstance(nota, list) else [nota], 6, normal=("Montserrat", BLUE),
                      destaque=("Montserrat SemiBold", BLUE))
        return s

    def sintese(self, titulo, esquerda, direita):
        """esquerda/direita: listas de (rótulo, texto) — rótulo em SemiBold azul."""
        s = self._clone("sintese")
        self._titulo(s, titulo)
        for nome, blocos in (("coluna_esq", esquerda), ("coluna_dir", direita)):
            preencher(self._shape(s, nome), [f"**{r}** {t}" for r, t in blocos], 10.5)
        return s

    def encerramento(self):
        return self._clone("encerramento")

    def salvar(self, caminho):
        lst = self.prs.slides._sldIdLst
        for sld in list(lst)[:len(self._base)]:
            self.prs.part.drop_rel(sld.get(f"{{{R_NS}}}id"))
            lst.remove(sld)
        self._base = []
        Path(caminho).parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(caminho))
        return caminho

    # ── gráficos nativos ──────────────────────────────────────────────────────
    def _add_chart(self, slide, g, x, y, w, h):
        tipo = g["tipo"]
        fmt = g.get("formato", "#,##0.00")
        cd = CategoryChartData(number_format=fmt)
        cd.categories = g["categorias"]
        for nome, vals in g["series"].items():
            cd.add_series(str(nome), vals)
        xl = {"colunas": XL_CHART_TYPE.COLUMN_CLUSTERED, "barras": XL_CHART_TYPE.BAR_CLUSTERED,
              "linhas": XL_CHART_TYPE.LINE_MARKERS}[tipo]
        gf = slide.shapes.add_chart(xl, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        gf.name = "grafico"
        ch = gf.chart
        n = len(g["series"])
        cores = SERIES.get(n, [GRAY] * (n - 2) + [BLUE_L, BLUE])
        rotular = g.get("rotular", {"colunas": "todas", "barras": "todas", "linhas": "fim"}[tipo])

        ch.font.name, ch.font.size = "Montserrat", Pt(8)
        ch.font.color.rgb = RGBColor.from_string(INK)
        ch.has_legend = True
        ch.legend.include_in_layout = False
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM if tipo == "linhas" else XL_LEGEND_POSITION.TOP
        ch.legend.font.size = Pt(8)

        cat, val = ch.category_axis, ch.value_axis
        cat.format.line.color.rgb = RGBColor.from_string(GRID)
        cat.has_major_gridlines = val.has_major_gridlines = False
        cat.tick_labels.font.size = Pt(7.5 if tipo == "barras" else 8)
        _tick_low(cat)

        plot = ch.plots[0]
        for i, ser in enumerate(plot.series):
            cor = RGBColor.from_string(cores[i])
            ultima = i == n - 1
            if tipo == "linhas":
                ser.format.line.color.rgb = cor
                ser.format.line.width = Pt(2.5 if ultima else 1.5)
                ser.smooth = False
                ser.marker.style = XL_MARKER_STYLE.CIRCLE
                ser.marker.size = 4
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = cor
                ser.marker.format.line.color.rgb = cor
                # rótulo só no último ponto: de todas as séries ('fim') ou só da corrente ('ultima')
                k = len(g["categorias"]) - 1
                if rotular == "fim" or ultima:
                    dl = ser.points[k].data_label
                    dl.has_text_frame = False
                    dl.position = XL_LABEL_POSITION.RIGHT
                    dl.font.size, dl.font.bold = Pt(8), True
                    dl.font.color.rgb = cor
                    _show_val(dl._get_or_add_dLbl())
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = cor
                ser.invert_if_negative = False
                if rotular == "todas" or ultima:
                    dls = ser.data_labels
                    dls.number_format, dls.number_format_is_linked = fmt, False
                    dls.position = XL_LABEL_POSITION.OUTSIDE_END
                    dls.show_value = True
                    dls.font.size = Pt(7 if tipo == "colunas" else 7.5)
                    dls.font.bold = ultima or tipo == "barras"
                    dls.font.color.rgb = RGBColor.from_string(BLUE if ultima else (INK if tipo == "colunas" else cores[i]))

        if tipo == "linhas":
            val.visible = True
            val.format.line.fill.background()
            val.tick_labels.font.size = Pt(8)
            val.tick_labels.number_format, val.tick_labels.number_format_is_linked = "#,##0", False
            if g.get("eixo_min") is not None:
                val.minimum_scale = g["eixo_min"]
            if g.get("eixo_max") is not None:
                val.maximum_scale = g["eixo_max"]
            _manual_layout(ch, 0.07, 0.05, 0.86, 0.74)
        else:
            val.visible = False
            plot.gap_width = 60 if tipo == "colunas" else 45
            plot.overlap = -5
            if tipo == "barras":
                cat.reverse_order = True   # 1ª categoria no topo
        return ch


PLOT_Y, PLOT_H = 0.13, 0.84     # área interna do gráfico de barras com tabela (frações)


def _tick_low(axis):
    ax = axis._element
    t = ax.find(qn("c:tickLblPos"))
    if t is None:
        t = etree.SubElement(ax, qn("c:tickLblPos"))
    t.set("val", "low")


def _show_val(dLbl):
    for tag in ("c:showLegendKey", "c:showVal", "c:showCatName", "c:showSerName", "c:showPercent", "c:showBubbleSize"):
        e = dLbl.find(qn(tag))
        if e is None:
            e = etree.SubElement(dLbl, qn(tag))
        e.set("val", "1" if tag == "c:showVal" else "0")


def _manual_layout(chart, x, y, w, h):
    plot = chart._chartSpace.find(f".//{{{C_NS}}}plotArea")
    lay = plot.find(f"{{{C_NS}}}layout")
    if lay is None:
        lay = etree.Element(f"{{{C_NS}}}layout")
        plot.insert(0, lay)
    for ch in list(lay):
        lay.remove(ch)
    ml = etree.SubElement(lay, f"{{{C_NS}}}manualLayout")
    for tag, v in [("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"),
                   ("x", x), ("y", y), ("w", w), ("h", h)]:
        etree.SubElement(ml, f"{{{C_NS}}}{tag}").set("val", str(v))


def _borda(cell, lado, cor=None, w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn(f"a:ln{lado}"))
    if old is not None:
        tcPr.remove(old)
    ln = etree.Element(qn(f"a:ln{lado}"), w=str(w if cor else 0))
    if cor:
        etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr"), val=cor)
    else:
        etree.SubElement(ln, qn("a:noFill"))
    ordem = ["L", "R", "T", "B"]
    pos = sum(1 for t in ordem[:ordem.index(lado)] if tcPr.find(qn(f"a:ln{t}")) is not None)
    tcPr.insert(pos, ln)


def _cel(cell, texto, cor, bold, size=7.5):
    cell.fill.background()
    cell.margin_left = cell.margin_right = Inches(0.04)
    cell.margin_top = cell.margin_bottom = 0
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.text_frame.word_wrap = False
    p = cell.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = texto
    r.font.size, r.font.bold = Pt(size), bold
    r.font.name = "Montserrat SemiBold" if bold else "Montserrat"
    r.font.color.rgb = RGBColor.from_string(cor)
    for lado in ("L", "R", "T"):
        _borda(cell, lado)


def _tabela_alinhada(slide, tabela, x, y, h, n, col_w=0.72, head_h=0.24):
    """Tabela de variação com uma linha por barra (mesma altura da faixa da categoria).
    Positivos em azul, negativos em tinta; colunas de texto (ex.: código CNAE) em tinta normal."""
    head, linhas = tabela["cabecalho"], tabela["linhas"]
    ncol = len(head)
    row_h = h * PLOT_H / n
    gf = slide.shapes.add_table(n + 1, ncol, Inches(x), Inches(y + h * PLOT_Y - head_h),
                                Inches(ncol * col_w), Inches(head_h + n * row_h))
    gf.name = "tabela_variacao"
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)
    for j in range(ncol):
        tbl.columns[j].width = Inches(col_w)
    tbl.rows[0].height = Inches(head_h)
    for j, t in enumerate(head):
        c = tbl.cell(0, j)
        _cel(c, t, BLUE, True)
        _borda(c, "B", BLUE, 9525)
    for i, lin in enumerate(linhas, start=1):
        tbl.rows[i].height = Inches(row_h)
        for j, t in enumerate(lin):
            c = tbl.cell(i, j)
            numerico = bool(re.match(r"^[+\-−]?\d", t)) and t[:1] in "+-−"
            cor = (INK if t.startswith(("-", "−")) else BLUE) if numerico else INK
            _cel(c, t, cor, numerico)
            _borda(c, "B", GRID if i < n else None)
    return gf
